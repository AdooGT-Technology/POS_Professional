from pathlib import Path
from datetime import date
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.pdfmetrics import stringWidth
from config import EXPORT_DIR

def export_excel(summary, top_products, inventory, d1, d2):
    path=EXPORT_DIR/f"sales_report_{d1}_{d2}.xlsx"
    wb=Workbook();ws=wb.active;ws.title="ملخص"
    ws.append(["تقرير المبيعات",f"{d1} إلى {d2}"])
    ws.append(["عدد الفواتير","الإجمالي قبل الخصم","الخصم","الضريبة","الإجمالي"])
    ws.append(list(summary))
    tp=wb.create_sheet("الأكثر مبيعاً");tp.append(["المنتج","الكمية","المبيعات"])
    for r in top_products:tp.append(list(r))
    inv=wb.create_sheet("المخزون");inv.append(["المنتج","المخزون","الحد الأدنى"])
    for r in inventory:inv.append(list(r))
    for sheet in wb.worksheets:
        for col in sheet.columns:
            letter=col[0].column_letter
            maxlen=max(len(str(x.value or "")) for x in col)
            sheet.column_dimensions[letter].width=min(maxlen+3,45)
        sheet.freeze_panes="A2"
    wb.save(path);return path

def _font():
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/tahoma.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont("POSUnicode", path))
                return "POSUnicode"
            except Exception:
                pass
    return "Helvetica"

def export_pdf(summary, top_products, inventory, d1, d2, title="تقرير المبيعات"):
    path=EXPORT_DIR/f"sales_report_{d1}_{d2}.pdf"
    c=canvas.Canvas(str(path),pagesize=A4);w,h=A4
    font=_font()
    y=h-45
    def line(txt,size=11):
        nonlocal y
        c.setFont(font,size);c.drawRightString(w-45,y,str(txt)[:110]);y-=18
        if y<55:c.showPage();y=h-45
    line(title,16);line(f"الفترة: {d1} - {d2}",11);line("")
    labels=["عدد الفواتير","قبل الخصم","الخصم","الضريبة","الإجمالي"]
    for lab,val in zip(labels,summary):line(f"{lab}: {val}")
    line("");line("الأكثر مبيعاً",14)
    for row in top_products:line(f"{row[0]} | كمية: {row[1]} | مبيعات: {row[2]}")
    line("");line("المخزون",14)
    for row in inventory:line(f"{row[0]} | مخزون: {row[1]} | حد أدنى: {row[2]}")
    c.save();return path
